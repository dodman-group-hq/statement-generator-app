"""
Unit tests for logic that needs neither a database nor LibreOffice:
password/session handling, plan defaults, StatementGenerator's header and
template validation, and the filesystem path guard.

Run with:  pytest
"""
import openpyxl
import pytest
from docx import Document

from app.accounts import (
    hash_password, verify_password, make_session_token, decode_session_token,
)
from app.db import plan_defaults, hash_key
from app.routes.generate_statements import StatementGenerator, _safe_join, _client_root


# ---------------------------------------------------------------------------
# Passwords and sessions
# ---------------------------------------------------------------------------

def test_password_hash_roundtrip():
    hashed = hash_password("correct horse battery")
    assert hashed != "correct horse battery"
    assert verify_password("correct horse battery", hashed)
    assert not verify_password("wrong password", hashed)


def test_verify_password_with_garbage_hash_is_false_not_an_exception():
    assert verify_password("anything", "not-a-bcrypt-hash") is False


def test_session_token_roundtrip():
    token = make_session_token("user-123")
    assert decode_session_token(token) == "user-123"


def test_tampered_session_token_is_rejected():
    token = make_session_token("user-123")
    user_id, nonce, sig = token.split(":")
    assert decode_session_token(f"other-user:{nonce}:{sig}") is None
    assert decode_session_token(f"{user_id}:{nonce}:{'0' * len(sig)}") is None


@pytest.mark.parametrize("bad", ["", "no-colons", "a:b", "a:b:c:d"])
def test_malformed_session_tokens_are_rejected(bad):
    assert decode_session_token(bad) is None


# ---------------------------------------------------------------------------
# Plan defaults and key hashing
# ---------------------------------------------------------------------------

def test_free_plan_defaults():
    assert plan_defaults("free") == {
        "monthly_job_limit": 5,
        "max_rows_per_job": 25,
        "password_protection_allowed": False,
    }


@pytest.mark.parametrize("plan", ["paid", "internal", "anything-else"])
def test_non_free_plans_are_unlimited_with_protection(plan):
    assert plan_defaults(plan) == {
        "monthly_job_limit": None,
        "max_rows_per_job": None,
        "password_protection_allowed": True,
    }


def test_plan_defaults_returns_a_copy():
    first = plan_defaults("free")
    first["monthly_job_limit"] = 9999
    assert plan_defaults("free")["monthly_job_limit"] == 5


def test_hash_key_is_deterministic_and_not_the_key():
    assert hash_key("sg_live_abc") == hash_key("sg_live_abc")
    assert hash_key("sg_live_abc") != hash_key("sg_live_abd")
    assert "sg_live_abc" not in hash_key("sg_live_abc")


# ---------------------------------------------------------------------------
# StatementGenerator: header detection and template validation
# ---------------------------------------------------------------------------

def _make_workbook(tmp_path, header, rows=()):
    path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return str(path)


def _make_template(tmp_path, text):
    path = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return str(path)


def _generator(tmp_path, header, template_text="Hello {{ name }}", rows=()):
    return StatementGenerator(
        template_path=_make_template(tmp_path, template_text),
        data_path=_make_workbook(tmp_path, header, rows),
        output_folder=str(tmp_path / "out"),
        temp_id="test-job",
    )


def test_generator_finds_id_and_name_columns_case_insensitively(tmp_path):
    gen = _generator(tmp_path, ["Member Name", "Member ID", "Balance"])
    assert gen.name_column_index == 0
    assert gen.id_column_index == 1


def test_generator_requires_an_id_column(tmp_path):
    with pytest.raises(ValueError, match="ID"):
        _generator(tmp_path, ["name", "balance"])


def test_generator_requires_a_name_column(tmp_path):
    with pytest.raises(ValueError, match="name"):
        _generator(tmp_path, ["id", "balance"])


def test_validate_template_passes_when_all_placeholders_have_columns(tmp_path):
    gen = _generator(tmp_path, ["id", "name", "Full Name"], template_text="{{ name }} / {{ Full_Name }}")
    gen.validate_template()  # header "Full Name" must satisfy {{ Full_Name }}


def test_validate_template_lists_missing_fields(tmp_path):
    gen = _generator(tmp_path, ["id", "name"], template_text="{{ name }} owes {{ balance }} by {{ due_date }}")
    with pytest.raises(ValueError) as exc:
        gen.validate_template()
    message = str(exc.value)
    assert "balance" in message and "due_date" in message
    assert "name," not in message  # 'name' exists, so it must not be reported


@pytest.mark.parametrize("value, expected", [
    (0, "-"),
    (None, "-"),
    (1234567, "1,234,567"),
    (1234.56, "1,235"),
    ("plain text", "plain text"),
])
def test_format_number(tmp_path, value, expected):
    gen = _generator(tmp_path, ["id", "name"])
    assert gen.format_number(value) == expected


def test_safe_id_from_filename(tmp_path):
    gen = _generator(tmp_path, ["id", "name"])
    assert gen._safe_id_from_filename("output_12345.docx") == "12345"
    assert gen._safe_id_from_filename("output_row2_Jane_Doe.pdf") == "row2_Jane_Doe"
    assert gen._safe_id_from_filename("something_else.pdf") is None


# ---------------------------------------------------------------------------
# Filesystem path guard
# ---------------------------------------------------------------------------

def test_safe_join_allows_paths_inside_the_root(tmp_path):
    root = str(tmp_path)
    assert _safe_join(root, "job-1", "file.pdf").startswith(str(tmp_path.resolve()))


@pytest.mark.parametrize("escape", ["../outside", "job/../../outside", "/etc/passwd"])
def test_safe_join_refuses_paths_outside_the_root(tmp_path, escape):
    with pytest.raises(ValueError):
        _safe_join(str(tmp_path), escape)


def test_client_roots_are_distinct_per_client():
    assert _client_root(1) != _client_root(2)
    assert _client_root(1).endswith("client_1")